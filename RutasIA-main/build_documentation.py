from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "documentacion"
OUT.mkdir(exist_ok=True)
LOGO = ROOT / "static" / "logistics" / "brand" / "logistiroute-logo.png"
TODAY = "17 de octubre de 2026"
CAPTURES = OUT / "capturas"
DIAGRAMS = OUT / "diagramas"

NAVY = "17324D"
BLUE = "276B8E"
PALE = "EAF3F7"
LIGHT = "F4F6F7"
GRAY = "5F6B73"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def configure(doc, short_title):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("20282D")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for sname, size, before, after in (("Title", 28, 0, 12), ("Heading 1", 19, 16, 7), ("Heading 2", 14, 12, 5), ("Heading 3", 11.5, 9, 3)):
        st = styles[sname]
        st.font.name = "Aptos Display"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header = sec.header
    hp = header.paragraphs[0]
    hp.text = short_title
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.style = styles["Normal"]
    for r in hp.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(GRAY)
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Flamia  |  ")
    r.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)
    return doc


def cover(doc, title, subtitle, audience, document_type):
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run().add_picture(str(LOGO), width=Inches(1.55))
    doc.add_paragraph(title, style="Title")
    p = doc.add_paragraph(subtitle)
    p.style = doc.styles["Subtitle"]
    p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.add_run("Sistema: ").bold = True
    p.add_run("Flamia - gestión de rutas, entregas y cobros")
    p = doc.add_paragraph()
    p.add_run("Documento: ").bold = True
    p.add_run(document_type)
    p = doc.add_paragraph()
    p.add_run("Dirigido a: ").bold = True
    p.add_run(audience)
    p = doc.add_paragraph()
    p.add_run("Versión documental: ").bold = True
    p.add_run("1.0")
    p = doc.add_paragraph()
    p.add_run("Fecha: ").bold = True
    p.add_run(TODAY)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.add_run("Documento de apoyo al proyecto de tesis").italic = True
    doc.add_page_break()


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def bullets(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        if isinstance(item, tuple):
            p.add_run(item[0]).bold = True
            p.add_run(item[1])
        else:
            p.add_run(item)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    hdr = t.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        c.text = str(h)
        set_cell_shading(c, NAVY)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(c)
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            if ridx % 2:
                set_cell_shading(cells[i], LIGHT)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.7)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def note(doc, lead, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.add_run(lead + " ").bold = True
    p.add_run(text)


def document_control(doc, purpose, scope):
    heading(doc, "Control documental", 1)
    table(doc, ["Elemento", "Descripción"], [
        ["Título", doc.core_properties.title], ["Sistema", "Flamia"], ["Versión", "1.0"],
        ["Fecha de emisión", TODAY], ["Propósito", purpose], ["Alcance", scope],
        ["Estado", "Documento final para revisión académica"],
    ], [1.7, 4.9])
    heading(doc, "Historial de revisiones", 2)
    table(doc, ["Versión", "Fecha", "Descripción"], [
        ["1.0", TODAY, "Emisión inicial basada en la revisión funcional y técnica del repositorio."],
    ], [1.0, 1.6, 4.0])
    doc.add_page_break()


def figure(doc, image_path, caption, width=6.75):
    path = Path(image_path)
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(10)
    for run in cap.runs:
        run.font.size = Pt(8.5)
        run.font.italic = True
        run.font.color.rgb = RGBColor.from_string(GRAY)


def make_diagrams():
    DIAGRAMS.mkdir(parents=True, exist_ok=True)
    font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 23)
    bold = ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf", 30)

    def box(draw, xy, title, subtitle, fill):
        draw.rounded_rectangle(xy, radius=20, fill=fill, outline="#B9C5CC", width=3)
        x1, y1, x2, _ = xy
        draw.text(((x1+x2)/2, y1+35), title, font=bold, fill="#142033", anchor="ma")
        draw.multiline_text(((x1+x2)/2, y1+88), subtitle, font=font, fill="#273647", anchor="ma", align="center", spacing=6)

    img = Image.new("RGB", (1600, 900), "white")
    d = ImageDraw.Draw(img)
    d.text((800, 45), "Arquitectura lógica de Flamia", font=bold, fill="#111111", anchor="ma")
    boxes = [
        ((100,150,1500,280), "Presentación", "Plantillas HTML  |  JavaScript  |  CSS  |  API JSON", "#EAF3F7"),
        ((100,350,1500,480), "Aplicación", "Planificación de viajes  |  Ciclo de vida  |  Transacciones", "#EEF2F5"),
        ((100,550,760,700), "Dominio", "Dijkstra  |  A estrella\nHaversine  |  Reglas", "#F4F6F7"),
        ((840,550,1500,700), "Infraestructura", "Django ORM  |  PostgreSQL\nMapas  |  Combustible", "#F4F6F7"),
    ]
    for b in boxes:
        box(d, *b)
    for y1, y2 in ((280,350),(480,550)):
        d.line((800,y1+10,800,y2-10), fill="#276B8E", width=8)
        d.polygon([(785,y2-25),(815,y2-25),(800,y2-5)], fill="#276B8E")
    img.save(DIAGRAMS / "arquitectura.png")

    img = Image.new("RGB", (1600, 900), "white")
    d = ImageDraw.Draw(img)
    d.text((800, 45), "Flujo operativo principal", font=bold, fill="#111111", anchor="ma")
    labels = ["Datos\nmaestros", "Entrega\npendiente", "Planificación", "Ruta en\nprogreso", "Parada\ny cobro", "Reporte"]
    xs = [70, 330, 590, 850, 1110, 1370]
    for i, (x, label) in enumerate(zip(xs, labels)):
        d.rounded_rectangle((x,300,x+170,520), radius=18, fill="#EAF3F7" if i%2==0 else "#F4F6F7", outline="#B9C5CC", width=3)
        d.multiline_text((x+85,410), label, font=font, fill="#142033", anchor="mm", align="center", spacing=8)
        if i < len(xs)-1:
            d.line((x+175,410,xs[i+1]-10,410), fill="#276B8E", width=7)
            d.polygon([(xs[i+1]-25,395),(xs[i+1]-25,425),(xs[i+1]-5,410)], fill="#276B8E")
    img.save(DIAGRAMS / "flujo_operativo.png")


def add_toc(doc, entries):
    heading(doc, "Contenido", 1)
    for n, title in entries:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(f"{n}  ").bold = True
        p.add_run(title)
    doc.add_page_break()


def build_user_manual():
    doc = configure(Document(), "Manual de usuario")
    doc.core_properties.title = "Manual de usuario de Flamia"
    doc.core_properties.subject = "Operación funcional del sistema de rutas entregas y cobros"
    doc.core_properties.author = "Proyecto Flamia"
    cover(doc, "Manual de usuario de Flamia", "Guía visual paso a paso para operar el sistema desde cero", "Administradores, supervisores, operadores y personal de reparto", "Manual de usuario con evidencia visual")
    document_control(doc, "Orientar a usuarios sin experiencia previa en el uso correcto del sistema.", "Ingreso, datos maestros, entregas, cobros, planificación, rutas, reportes, mapa y administración.")
    entries = [
        (1, "Antes de comenzar"), (2, "Ingreso y orientación"), (3, "Flujo recomendado de trabajo"),
        (4, "Control diario"), (5, "Vehículos"), (6, "Tiendas"), (7, "Productos"),
        (8, "Repartidores"), (9, "Entregas"), (10, "Cobros"), (11, "Planificador"),
        (12, "Rutas"), (13, "Reportes"), (14, "Mapa"), (15, "Zonas y usuarios"),
        (16, "Trabajo sin conexión"), (17, "Problemas frecuentes"), (18, "Glosario y listas de verificación")]
    add_toc(doc, entries)

    heading(doc, "1 Antes de comenzar")
    para(doc, "Flamia organiza tiendas, productos, entregas, cobros, vehículos y rutas de reparto en una sola pantalla. Este manual explica qué hacer, en qué orden y cómo reconocer que cada operación terminó correctamente. No se requieren conocimientos técnicos.")
    heading(doc, "Qué necesita", 2)
    bullets(doc, ["Una computadora, tableta o teléfono con un navegador actualizado.", "La dirección web del sistema proporcionada por su administrador.", "Un usuario y contraseña activos.", "Conexión a Internet para ingresar y cargar la información inicial. Los cobros pueden guardarse temporalmente cuando la conexión se interrumpe."])
    heading(doc, "Roles y permisos", 2)
    table(doc, ["Rol", "Puede hacer", "Limitaciones principales"], [
        ["Administrador", "Operar todo el sistema; crear usuarios y zonas; editar precio de combustible.", "Debe proteger sus credenciales y revisar datos maestros."],
        ["Supervisor", "Planificar, iniciar, completar o cancelar rutas y realizar la operación diaria.", "No crea usuarios ni modifica las funciones exclusivas de administrador."],
        ["Operador", "Consultar y registrar la operación permitida, incluidos pedidos, cobros y eventos.", "No ejecuta acciones reservadas a administrador o supervisor."],
    ], [1.1, 3.0, 2.5])
    note(doc, "Importante.", "Las opciones Zonas y Usuarios solamente aparecen para un administrador. Algunas acciones de rutas también dependen del rol.")

    heading(doc, "2 Ingreso y orientación")
    figure(doc, CAPTURES / "01_inicio_sesion.png", "Figura 1. Pantalla de inicio de sesión del sistema Flamia.")
    heading(doc, "Iniciar sesión", 2)
    bullets(doc, ["Abra el enlace de Flamia.", "Escriba su nombre de usuario en Usuario.", "Escriba su contraseña. Los caracteres se ocultan por seguridad.", "Pulse Ingresar.", "Confirme que en la parte superior aparece su nombre o usuario y su rol."], numbered=True)
    heading(doc, "Si no puede entrar", 2)
    bullets(doc, ["Revise mayúsculas, minúsculas y espacios accidentales.", "Confirme que utiliza el usuario correcto y que tiene conexión.", "Solicite al administrador que verifique que la cuenta esté activa.", "No comparta contraseñas ni las envíe por mensajes sin protección."])
    heading(doc, "Cómo está organizada la pantalla", 2)
    table(doc, ["Zona", "Uso"], [
        ["Menú superior", "Cambia entre Control diario, Vehículos, Tiendas, Productos, Repartidores, Entregas, Cobros, Planificador, Rutas, Reportes y Mapa."],
        ["Formularios", "Permiten registrar información. Los campos obligatorios deben completarse antes de guardar."],
        ["Listados", "Muestran registros ya creados, estados y botones disponibles."],
        ["Mensajes breves", "Confirman un guardado o explican un error. Léalo antes de repetir la operación."],
    ], [1.5, 5.1])

    heading(doc, "3 Flujo recomendado de trabajo")
    para(doc, "Para evitar listas vacías o rutas imposibles, prepare los datos en este orden:")
    bullets(doc, ["El administrador confirma zonas y usuarios.", "Registre repartidores y vehículos.", "Registre tiendas con ubicación correcta.", "Registre productos, precios y pesos.", "Cree las entregas pendientes.", "Planifique una ruta con un vehículo cuya capacidad sea suficiente.", "Inicie la ruta, actualice paradas y registre cobros o eventos.", "Complete la ruta y revise reportes y saldos."], numbered=True)
    note(doc, "Regla clave.", "Una entrega solo aparece para planificación cuando está Pendiente. Al planificar pasa a Asignado; al iniciar la ruta pasa a En tránsito; al finalizar queda Entregado.")
    figure(doc, DIAGRAMS / "flujo_operativo.png", "Figura 2. Secuencia general de trabajo en Flamia.")

    heading(doc, "4 Control diario")
    figure(doc, CAPTURES / "02_control_diario.png", "Figura 3. Tablero de control diario con filtros e indicadores operativos.")
    para(doc, "Es la página inicial y resume el estado de la operación. Puede filtrar por fechas y pulsar Actualizar para renovar los indicadores.")
    table(doc, ["Indicador", "Qué significa"], [
        ["Rutas planificadas", "Viajes preparados que todavía no terminaron."], ["Rutas activas", "Rutas planificadas o en progreso."],
        ["Entregas pendientes", "Pedidos todavía disponibles para planificar."], ["Entregas completadas", "Pedidos marcados como entregados."],
        ["Cobros estimados", "Costo estimado acumulado de los viajes, según el tablero."], ["Distancia acumulada", "Kilómetros estimados de las rutas."],
        ["Gráficas", "Distribución por estado, actividad de siete días y vehículos más utilizados."],
    ], [2.1, 4.5])
    heading(doc, "Precio de combustible", 2)
    para(doc, "El sistema muestra los precios de regular, súper y diésel y la fuente. Solo el administrador puede editar estos valores. El planificador utiliza el precio de gasolina regular para calcular el costo estimado de combustible.")

    heading(doc, "5 Vehículos")
    figure(doc, CAPTURES / "03_vehiculos.png", "Figura 4. Registro y consulta de vehículos de reparto.")
    para(doc, "Use esta opción para registrar la flota. Complete placa, modelo, capacidad en kilogramos, eficiencia en kilómetros por litro y costo por kilómetro. Opcionalmente asigne repartidor y zona base, deje Activo marcado y pulse Guardar vehículo.")
    table(doc, ["Campo", "Cómo llenarlo", "Ejemplo"], [
        ["Placa", "Identificador único; no repita una placa.", "C-123ABC"], ["Modelo", "Tipo o modelo reconocible.", "Isuzu NPR"],
        ["Capacidad", "Carga máxima segura en kg.", "2500"], ["Eficiencia", "Rendimiento aproximado km/l.", "6.20"],
        ["Costo por km", "Costo operativo usado en la estimación.", "4.10"], ["Activo", "Desmárquelo si no debe usarse para nuevas rutas.", "Marcado"],
    ], [1.4, 3.5, 1.5])
    note(doc, "Evite errores.", "La suma del peso de las entregas seleccionadas no puede superar la capacidad del vehículo y la eficiencia debe ser mayor que cero.")

    heading(doc, "6 Tiendas")
    figure(doc, CAPTURES / "04_tiendas.png", "Figura 5. Formulario de tiendas y datos de geolocalización.")
    para(doc, "Registre cada cliente o punto de entrega. Los datos de ubicación son esenciales para el mapa y el orden de las paradas.")
    bullets(doc, ["Escriba nombre, encargado y teléfono.", "Seleccione zona o departamento e indique municipio y colonia.", "Elija el producto habitual: agua, gas o ambos.", "Agregue dirección y una referencia fácil de reconocer.", "Ingrese latitud y longitud. Puede usar la ubicación del dispositivo si el navegador solicita permiso.", "Mantenga Activo marcado y guarde."], numbered=True)
    note(doc, "Coordenadas.", "La latitud de Guatemala suele ser positiva y la longitud negativa. Una coordenada incorrecta colocará la tienda en otro lugar del mapa.")

    heading(doc, "7 Productos")
    figure(doc, CAPTURES / "05_productos.png", "Figura 6. Catálogo de productos, presentaciones, precios y pesos.")
    para(doc, "Cree el catálogo que luego se usará en Entregas. Indique nombre, tipo, presentación, precio unitario y peso unitario. El total y la carga del pedido se calculan a partir de la cantidad.")
    table(doc, ["Dato", "Ejemplo"], [["Nombre", "Garrafón de agua"], ["Tipo", "Agua purificada"], ["Presentación", "5 galones"], ["Precio unitario", "Q18.00"], ["Peso unitario", "20 kg"]], [2.2, 4.4])

    heading(doc, "8 Repartidores")
    figure(doc, CAPTURES / "06_repartidores.png", "Figura 7. Registro y listado de repartidores.")
    para(doc, "Registre nombre completo, teléfono y número de licencia o identificación. El número debe ser único. Deje Activo marcado para que aparezca al asignar vehículos o planificar rutas.")

    heading(doc, "9 Entregas")
    figure(doc, CAPTURES / "07_entregas.png", "Figura 8. Creación, filtros y estados de las entregas.")
    para(doc, "Una entrega representa un pedido para una tienda o destino. El sistema genera automáticamente un código con formato PED fecha número.")
    bullets(doc, ["Seleccione una tienda registrada, si corresponde.", "Seleccione la base de salida y el destino.", "Elija el producto y escriba la cantidad.", "Revise la carga y el total estimados que aparecen en pantalla.", "Elija prioridad Baja, Normal o Alta.", "Pulse Crear entrega y confirme que aparece en la lista."], numbered=True)
    table(doc, ["Estado", "Interpretación"], [["Pendiente", "Disponible para planificar."], ["Asignado", "Incluido en una ruta planificada."], ["En ruta", "El viaje ya comenzó."], ["Entregado", "La entrega terminó."], ["Cancelado", "El pedido no continuará."]], [1.6, 5.0])
    para(doc, "Los filtros permiten buscar por estado, origen y destino. Úselos antes de concluir que una entrega desapareció.")

    heading(doc, "10 Cobros")
    figure(doc, CAPTURES / "08_cobros.png", "Figura 9. Registro de pagos, saldos por tienda e historial de cobros.")
    para(doc, "Seleccione tienda, pedido si aplica, monto, tipo de pago, método y una nota útil. Pulse Registrar cobro. El saldo por tienda se calcula comparando las ventas de sus pedidos con los pagos registrados.")
    table(doc, ["Opción", "Cuándo usarla"], [["Pago completo", "El cliente cancela el total."], ["Abono", "Entrega una parte de la deuda."], ["Crédito", "Se registra la operación como crédito."], ["Efectivo", "Dinero recibido físicamente."], ["Transferencia", "Pago bancario o electrónico."], ["Otro", "Un método distinto; explíquelo en la nota."]], [1.8, 4.8])
    note(doc, "Comprobación.", "Después de guardar, revise Historial de cobros y Saldos por tienda. Si estaba sin conexión, revise el indicador de sincronización al recuperar Internet.")

    heading(doc, "11 Planificador")
    figure(doc, CAPTURES / "09_planificador.png", "Figura 10. Selección de vehículo, repartidor y entregas para planificar una ruta.")
    para(doc, "El planificador combina entregas pendientes en una ruta. El sistema calcula el camino, distancia, consumo, costo de combustible y costo operativo; además crea las paradas en orden.")
    bullets(doc, ["Seleccione un vehículo activo.", "Seleccione un repartidor si no está asignado al vehículo.", "Marque una o más entregas pendientes.", "Revise que el peso total no exceda la capacidad.", "Pulse Planificar ruta.", "Lea el resultado y anote el código VIA generado."], numbered=True)
    note(doc, "Si falla.", "Compruebe que las entregas estén pendientes, el vehículo esté activo, exista conexión entre las zonas y la capacidad sea suficiente.")

    heading(doc, "12 Rutas")
    figure(doc, CAPTURES / "10_rutas.png", "Figura 11. Seguimiento y acciones disponibles para las rutas de reparto.")
    para(doc, "Aquí se consulta y controla cada viaje. Puede filtrar por estado, vehículo y fechas.")
    table(doc, ["Acción", "Efecto"], [
        ["Iniciar", "Cambia el viaje a En progreso y sus pedidos a En tránsito."],
        ["Completar", "Marca el viaje y, salvo paradas no entregadas, las entregas como completadas."],
        ["Cancelar", "Cancela el viaje y devuelve a Pendiente las entregas que aún no se entregaron."],
        ["Registrar evento", "Agrega una nota cronológica como visita, incidencia o cobro."],
        ["Actualizar parada", "Permite marcar Pendiente, Entregado o No entregado y añadir observación."],
    ], [1.6, 5.0])
    note(doc, "Precaución.", "No complete una ruta antes de revisar las paradas. Completar actualiza masivamente los pedidos asociados.")

    heading(doc, "13 Reportes")
    figure(doc, CAPTURES / "11_reportes.png", "Figura 12. Resumen operativo y financiero filtrado por fechas.")
    para(doc, "Seleccione Desde y Hasta y actualice el resumen. Se muestran ventas, pagado, saldo, cantidad de pedidos, rutas, distancia y costo estimado de combustible. Los filtros usan la fecha de creación de los registros.")
    bullets(doc, ["Ventas: suma de productos incluidos en pedidos.", "Pagado: suma de cobros registrados.", "Saldo: diferencia no negativa entre ventas y pagos.", "Combustible: suma del costo estimado almacenado en los viajes."])

    heading(doc, "14 Mapa")
    figure(doc, CAPTURES / "12_mapa.png", "Figura 13. Mapa de rutas, puntos de interés y simulación de recorrido.")
    para(doc, "El mapa muestra zonas, tiendas, paradas y rutas. En Resaltar ruta del día seleccione un viaje para ver su trayecto y resumen. La geometría por carretera depende del servicio de mapas configurado; si no está disponible, el sistema puede mostrar una línea aproximada entre puntos.")
    bullets(doc, ["Verifique que las tiendas tengan coordenadas correctas.", "Seleccione una ruta y espere a que cargue el trazado.", "Pulse o toque marcadores para consultar detalles.", "La simulación GPS es visual; no sustituye el seguimiento real del vehículo."])

    heading(doc, "15 Zonas y usuarios")
    heading(doc, "Zonas", 2)
    para(doc, "Solo el administrador puede ver esta sección. Puede elegir un departamento de Guatemala para autocompletar o escribir código, nombre, latitud y longitud. Las conexiones entre zonas que usa el algoritmo se administran desde el panel administrativo de Django.")
    figure(doc, CAPTURES / "13_zonas.png", "Figura 14. Administración de zonas y coordenadas geográficas.")
    heading(doc, "Usuarios", 2)
    para(doc, "Solo el administrador crea cuentas. Complete usuario, contraseña, nombre, apellido, correo y rol. Entregue las credenciales por un medio seguro y pida al usuario que no las comparta.")
    figure(doc, CAPTURES / "14_usuarios.png", "Figura 15. Creación de usuarios y asignación de roles.")

    heading(doc, "16 Trabajo sin conexión")
    para(doc, "Flamia instala componentes de aplicación web progresiva y conserva recursos básicos del sitio. Cuando falla la red, el registro de ciertos cobros puede añadirse a una cola local y sincronizarse después.")
    bullets(doc, ["No cierre ni borre los datos del navegador si existen registros pendientes.", "Mantenga abierta la aplicación hasta recuperar conexión cuando sea posible.", "Observe el estado de sincronización y confirme que el cobro aparece en el historial.", "Evite registrar el mismo cobro desde dos dispositivos."])
    note(doc, "Alcance.", "El modo sin conexión no garantiza que todas las pantallas o consultas funcionen sin Internet; está orientado principalmente a conservar y sincronizar operaciones compatibles.")

    heading(doc, "17 Problemas frecuentes")
    table(doc, ["Problema", "Qué revisar"], [
        ["No puedo iniciar sesión", "Usuario, contraseña, cuenta activa, dirección correcta y conexión."],
        ["No aparece una opción", "Rol asignado; Zonas y Usuarios son exclusivos del administrador."],
        ["No aparece una entrega", "Filtros activos y estado Pendiente."],
        ["La ruta no se crea", "Capacidad, vehículo activo, pedidos pendientes y conexión entre zonas."],
        ["La tienda aparece lejos", "Latitud y longitud, especialmente el signo negativo de longitud."],
        ["El saldo parece incorrecto", "Productos y cantidades del pedido, cobros duplicados y tienda/pedido asociados."],
        ["El mapa no dibuja carretera", "Conexión a Internet y clave del proveedor; el trazo aproximado puede seguir disponible."],
        ["Un cobro está pendiente", "Recupere Internet y permita la sincronización antes de volver a registrarlo."],
    ], [2.2, 4.4])

    heading(doc, "18 Glosario y listas de verificación")
    table(doc, ["Término", "Significado sencillo"], [["Entrega o pedido", "Solicitud de productos para un destino."], ["Ruta o viaje", "Recorrido asignado a un vehículo con una o más entregas."], ["Parada", "Tienda o destino dentro de una ruta."], ["Carga", "Peso total transportado."], ["Eficiencia", "Kilómetros que el vehículo recorre por litro."], ["Sincronizar", "Enviar al servidor un registro guardado temporalmente en el dispositivo."], ["Dijkstra y A estrella", "Métodos matemáticos para encontrar un camino corto en la red de zonas."]], [2.0, 4.6])
    heading(doc, "Antes de salir a reparto", 2)
    bullets(doc, ["Vehículo y repartidor activos.", "Capacidad suficiente.", "Tiendas y coordenadas verificadas.", "Entregas correctas y pendientes.", "Ruta planificada y visible en el mapa.", "Precio de combustible actualizado."])
    heading(doc, "Al terminar el día", 2)
    bullets(doc, ["Paradas actualizadas.", "Cobros sincronizados.", "Incidencias registradas.", "Rutas completadas o canceladas correctamente.", "Saldos y reporte revisados.", "Sesión cerrada en equipos compartidos."])
    doc.save(OUT / "Manual_de_Usuario_Flamia.docx")


def build_technical():
    doc = configure(Document(), "Documentación técnica")
    doc.core_properties.title = "Documentación técnica de Flamia"
    doc.core_properties.subject = "Arquitectura desarrollo despliegue y mantenimiento"
    doc.core_properties.author = "Proyecto Flamia"
    cover(doc, "Documentación técnica de Flamia", "Arquitectura, datos, API, despliegue y mantenimiento", "Desarrolladores, administradores de plataforma y personal de soporte", "Memoria técnica del sistema")
    document_control(doc, "Describir la solución implementada y facilitar su evaluación, despliegue y mantenimiento.", "Arquitectura, repositorio, configuración, datos, reglas, algoritmos, API, seguridad, pruebas y operación.")
    entries = [(1,"Resumen técnico"),(2,"Arquitectura"),(3,"Estructura del repositorio"),(4,"Configuración"),(5,"Modelo de datos"),(6,"Reglas de negocio"),(7,"Optimización de rutas"),(8,"API"),(9,"Cliente web y PWA"),(10,"Seguridad y permisos"),(11,"Instalación local"),(12,"Despliegue en Render"),(13,"Operación y mantenimiento"),(14,"Pruebas"),(15,"Riesgos y deuda técnica"),(16,"Referencia rápida")]
    add_toc(doc, entries)
    heading(doc, "1 Resumen técnico")
    para(doc, "Flamia es una aplicación monolítica Django 6 orientada a logística de reparto, ventas y cobros. Expone una interfaz de página única construida con HTML, CSS y JavaScript sin framework, y un conjunto de endpoints JSON protegidos por sesión. Persiste en SQLite por defecto o PostgreSQL cuando se configura DATABASE_URL.")
    table(doc, ["Componente", "Tecnología", "Responsabilidad"], [
        ["Servidor", "Python 3.12 y Django 6.0.4", "Autenticación, reglas, API, ORM y plantillas."],
        ["Producción", "Gunicorn 22", "Servidor WSGI con tres workers."], ["Datos", "SQLite o PostgreSQL", "Persistencia transaccional."],
        ["Estáticos", "WhiteNoise 6.7", "Archivos versionados y comprimidos."], ["Cliente", "JavaScript, HTML y CSS", "SPA, formularios, filtros, gráficas y mapa."],
        ["Mapas", "MapLibre, OpenFreeMap, ORS y Google Maps opcional", "Tiles, geometría y visualización."], ["Scraping", "httpx y BeautifulSoup", "Consulta programada de precios de combustible."],
    ], [1.4, 1.8, 3.4])

    heading(doc, "2 Arquitectura")
    para(doc, "El paquete logistics separa presentación, aplicación, dominio e infraestructura. Django ORM atraviesa varias capas, por lo que la separación es pragmática y no una implementación DDD estrictamente aislada.")
    figure(doc, DIAGRAMS / "arquitectura.png", "Figura 1. Arquitectura lógica y responsabilidades principales de Flamia.")
    table(doc, ["Capa", "Archivos", "Responsabilidad"], [
        ["Presentación", "presentation/views.py y serializers.py", "HTTP, permisos, validación básica, serialización JSON."],
        ["Aplicación", "application/services.py", "Planificación y ciclo de vida transaccional."],
        ["Dominio", "domain/services.py y exceptions.py", "Dijkstra, A estrella, Haversine y errores de planificación."],
        ["Infraestructura", "infrastructure/fuel_scraper.py", "Acceso externo para precios de combustible."],
        ["Persistencia", "models.py y migrations", "Esquema relacional y propiedades financieras."],
        ["Cliente", "templates y static", "Interfaz, estado del navegador, PWA y mapas."],
    ], [1.3, 2.4, 3.0])
    heading(doc, "Flujo de planificación", 2)
    bullets(doc, ["El navegador envía vehicle_id, order_ids, driver_id opcional y algoritmo.", "La vista valida rol de administrador o supervisor y carga objetos activos.", "TripPlanner valida estado de pedidos, capacidad y eficiencia.", "El optimizador construye el grafo desde RouteConnection y calcula cada tramo.", "Una transacción crea Trip, TripOrder, RouteStop y TripEvent.", "Los pedidos cambian de Pendiente a Asignado y se devuelve el viaje serializado."], numbered=True)
    figure(doc, DIAGRAMS / "flujo_operativo.png", "Figura 2. Flujo operativo desde la preparación de datos hasta el reporte.")

    heading(doc, "3 Estructura del repositorio")
    table(doc, ["Ruta", "Contenido"], [
        ["flamia/settings.py", "Configuración, base de datos, seguridad, idioma, zona horaria y estáticos."], ["flamia/urls.py", "Admin, autenticación, manifiesto, service worker y aplicación."],
        ["logistics/models.py", "Entidades y relaciones."], ["logistics/presentation", "API y serializadores."], ["logistics/application", "Casos de uso de rutas."],
        ["logistics/domain", "Algoritmos y excepción PlanningError."], ["logistics/infrastructure", "Obtención de combustible."], ["logistics/management/commands", "Carga demo y sincronización."],
        ["templates", "Login e interfaz principal."], ["static", "JavaScript, CSS, marca y PWA."], ["render.yaml", "Web service y cron job en Render."],
    ], [2.4, 4.3])

    heading(doc, "4 Configuración")
    table(doc, ["Variable", "Obligatoria", "Descripción"], [
        ["SECRET_KEY", "Producción", "Clave criptográfica de Django. No usar el valor inseguro por defecto."],
        ["DEBUG", "Sí", "False en producción; True solo en desarrollo."], ["DATABASE_URL", "Producción", "postgres:// o postgresql://. Si falta, usa db.sqlite3."],
        ["ALLOWED_HOSTS", "Según entorno", "Lista separada por comas; incluye localhost y .onrender.com por defecto."],
        ["CSRF_TRUSTED_ORIGINS", "Según dominio", "Orígenes HTTPS autorizados, separados por comas."],
        ["ORS_API_KEY", "Opcional", "Clave de OpenRouteService para geometría vial."], ["GOOGLE_MAPS_API_KEY", "Opcional", "Clave de Google Maps usada por el cliente cuando está disponible."],
    ], [1.9, 1.1, 3.7])
    para(doc, "La zona horaria es America/Guatemala, el idioma es es-gt, USE_TZ está activo y la sesión expira al cerrar el navegador o a las ocho horas. En producción, las cookies de sesión y CSRF se marcan Secure.")

    heading(doc, "5 Modelo de datos")
    table(doc, ["Entidad", "Campos clave", "Relaciones"], [
        ["UserProfile", "role", "Uno a uno con User."], ["Department", "code, name, latitude, longitude", "Origen/destino de conexiones, pedidos, viajes y vehículos."],
        ["RouteConnection", "distance_km, is_bidirectional", "Department origen y destino; combinación única."], ["Driver", "license_number, is_active", "Asignable a vehículos y viajes."],
        ["Vehicle", "plate, capacity_kg, efficiency, cost_per_km", "Zona actual y repartidor opcionales."], ["Store", "ubicación, product_type, is_active", "Pertenece a Department; pedidos, cobros y paradas."],
        ["Product", "type, presentation, price, weight", "Referenciado por OrderItem."], ["Order", "code, weight, priority, status", "Tienda opcional; origen/destino; ítems, pagos y viajes."],
        ["OrderItem", "quantity, unit_price, subtotal", "Order y Product."], ["Trip", "ruta, distancia, consumo, costos, status", "Vehicle, Driver, Departments y Orders mediante TripOrder."],
        ["RouteStop", "sequence, status, coordinates", "Trip, Order y Store opcionales."], ["Payment", "uuid, amount, type, method, sync_status", "Store obligatoria; Order y Trip opcionales."],
        ["TripEvent", "note, timestamps", "Pertenece a Trip."],
    ], [1.2, 2.7, 2.8])
    heading(doc, "Estados", 2)
    table(doc, ["Objeto", "Estados"], [["Order", "pending, assigned, in_transit, delivered, canceled"], ["Trip", "planned, in_progress, completed, canceled"], ["RouteStop", "pending, delivered, not_delivered"], ["Payment sync", "synced, pending"]], [1.6, 5.1])
    para(doc, "Los códigos PED y VIA combinan la fecha local con cuatro dígitos aleatorios. Son únicos en base de datos, pero no existe reintento explícito ante una colisión excepcional.")

    heading(doc, "6 Reglas de negocio")
    bullets(doc, ["Solo pedidos Pendiente pueden entrar a una planificación.", "El vehículo debe estar activo, tener eficiencia mayor que cero y capacidad suficiente.", "La planificación y los cambios de ciclo de vida se ejecutan dentro de transacciones atómicas.", "Planificar crea paradas ordenadas y cambia pedidos a Asignado.", "Iniciar requiere estado Planificado y cambia pedidos a En tránsito.", "Completar acepta Planificado o En progreso, completa pedidos y todas las paradas salvo No entregado.", "Cancelar no acepta un viaje Completado y devuelve pedidos no entregados a Pendiente.", "Marcar una parada Entregada actualiza su pedido a Entregado; No entregado lo devuelve a Pendiente.", "El total del pedido es la suma de subtotales. El saldo pendiente es el mayor valor entre total menos pagado y cero."])
    note(doc, "Observación de implementación.", "La validación de que todos los pedidos compartan origen y destino está desactivada de hecho por la expresión True or all(...). El planificador actual sí admite múltiples destinos y los ordena por cercanía de red.")

    heading(doc, "7 Optimización de rutas")
    para(doc, "RouteOptimizer implementa Dijkstra sobre un grafo de departamentos. AStarOptimizer añade una heurística Haversine cuando ambos nodos tienen coordenadas. Las distancias se manejan con Decimal y se redondean a dos decimales.")
    table(doc, ["Aspecto", "Dijkstra", "A estrella"], [["Prioridad", "Distancia acumulada g", "g más distancia Haversine h"], ["Óptimo", "Sí, con pesos no negativos", "Sí si la heurística no sobreestima la distancia del grafo"], ["Coordenadas", "No necesarias", "Si faltan, h vale cero y se aproxima a Dijkstra"], ["Uso del cliente", "Predeterminado y rotulado en la UI", "Aceptado por la API mediante algorithm=astar"]], [1.6, 2.5, 2.5])
    para(doc, "El grafo y las coordenadas se guardan en la caché Django durante 120 segundos. Existe invalidate_route_cache(), pero las vistas que crean Department o RouteConnection deben invocarla cuando corresponda. La ruta multi parada selecciona iterativamente la siguiente entrega con menor distancia de red desde el punto actual; es una heurística voraz y no garantiza el óptimo global del recorrido completo.")
    para(doc, "Consumo: litros = distancia / km_por_litro; galones = litros / 3.78541. Costo de combustible = galones por precio regular. Costo operativo = distancia por costo_por_km.")

    heading(doc, "8 API")
    para(doc, "Todos los endpoints requieren sesión autenticada y responden JSON con ok=true o ok=false y error. Las mutaciones usan POST y protección CSRF de Django.")
    table(doc, ["Ruta", "Método", "Función y permiso especial"], [
        ["/api/me/", "GET", "Usuario y rol."], ["/api/users/", "GET POST", "Listar o crear usuarios; solo admin."], ["/api/dashboard/", "GET", "Indicadores y series."],
        ["/api/departments/", "GET POST", "Listar; crear solo admin."], ["/api/connections/", "GET", "Conexiones del grafo."], ["/api/stores/", "GET POST", "Tiendas."],
        ["/api/products/", "GET POST", "Productos."], ["/api/drivers/", "GET POST", "Repartidores."], ["/api/vehicles/", "GET POST", "Vehículos."],
        ["/api/orders/", "GET POST", "Pedidos, filtros y creación."], ["/api/trips/", "GET", "Viajes filtrados."], ["/api/trips/plan/", "POST", "Planificar; admin o supervisor."],
        ["/api/trips/{id}/action/", "POST", "start, complete o cancel; admin o supervisor."], ["/api/trips/{id}/events/", "POST", "Agregar evento."],
        ["/api/trips/{id}/stops/{stop}/", "POST", "Actualizar estado y observación."], ["/api/payments/", "GET POST", "Cobros; UUID permite idempotencia."],
        ["/api/balances/", "GET", "Saldos agregados por tienda."], ["/api/reports/", "GET", "Resumen entre date_from y date_to."], ["/api/fuel-price/", "GET POST", "Consultar; editar solo admin."],
    ], [2.4, 0.9, 3.4])
    heading(doc, "Contratos relevantes", 2)
    bullets(doc, ["Planificación: vehicle_id, order_ids, driver_id opcional, algorithm opcional.", "Cobro: store_id, amount, order_id y trip_id opcionales, payment_type, method, sync_status, note y uuid opcional.", "Acción de viaje: action con start, complete o cancel.", "Parada: status y observation.", "Filtros: fechas ISO AAAA-MM-DD; órdenes por status, origin_id y destination_id; viajes por status, vehicle_id y fechas."])

    heading(doc, "9 Cliente web y PWA")
    para(doc, "index.html contiene todas las secciones; app.js mantiene el estado, consulta APIs, procesa formularios, renderiza tablas, dibuja gráficas en canvas y controla MapLibre. Las pestañas se enlazan mediante data-tab y pueden activarse desde la URL.")
    bullets(doc, ["El token CSRF se lee de la cookie csrftoken y se envía en POST.", "service-worker.js cachea el shell básico y facilita la apertura de la aplicación instalada.", "manifest.json define nombre, colores, iconos y modo standalone.", "Los cobros compatibles usan localStorage como cola cuando se detecta error de red; un UUID evita crear duplicados al sincronizar.", "La ruta vial intenta Google Maps si está configurado y utiliza servicios cartográficos según disponibilidad; el cliente conserva un trazado alternativo.", "La simulación GPS interpola posiciones para demostración visual y puede automatizar entregas en el cliente."])

    heading(doc, "10 Seguridad y permisos")
    bullets(doc, ["Autenticación por sesión de Django; login en /accounts/login/ y logout estándar.", "CSRF activo en middleware y cookie inicializada en la vista principal.", "Cookies Secure en producción y soporte de proxy HTTPS.", "Validadores de contraseña estándar de Django configurados.", "Permisos por rol aplicados explícitamente a usuarios, creación de zonas, precio de combustible, planificación y acciones de viaje.", "El panel /admin/ requiere permisos de staff de Django, independientes de UserProfile.role."])
    heading(doc, "Recomendaciones de endurecimiento", 2)
    bullets(doc, ["Exigir SECRET_KEY en producción y eliminar el fallback inseguro.", "Configurar ALLOWED_HOSTS y CSRF_TRUSTED_ORIGINS de forma explícita.", "Revisar autorización de escritura para tiendas, productos, repartidores, vehículos, pedidos, cobros, eventos y paradas; actualmente varios POST solo exigen autenticación.", "Evitar devolver str(exc) al cliente para excepciones imprevistas; registrar internamente y responder un mensaje genérico.", "Añadir límites, auditoría y rotación de credenciales para proveedores de mapas.", "Aplicar copias de seguridad y política de retención de datos financieros y personales."])

    heading(doc, "11 Instalación local")
    para(doc, "Requisitos: Python 3.12, pip y Git. PostgreSQL es opcional porque sin DATABASE_URL se usa SQLite.")
    bullets(doc, ["Entrar en la carpeta que contiene manage.py.", "Crear un entorno virtual: python -m venv .venv.", "Activarlo en PowerShell: .venv\\Scripts\\Activate.ps1.", "Instalar dependencias: pip install -r requirements.txt.", "Crear .env con DEBUG=True y, si aplica, claves y DATABASE_URL.", "Ejecutar python manage.py migrate.", "Cargar demostración con python manage.py seed_demo_data, si se desea.", "Crear administrador real con python manage.py createsuperuser.", "Ejecutar python manage.py runserver y abrir http://127.0.0.1:8000/."], numbered=True)
    heading(doc, "Comandos útiles", 2)
    table(doc, ["Comando", "Uso"], [["python manage.py check", "Valida configuración."], ["python manage.py makemigrations", "Genera migraciones tras cambiar modelos."], ["python manage.py migrate", "Aplica esquema."], ["python manage.py test", "Ejecuta pruebas."], ["python manage.py collectstatic --noinput", "Prepara estáticos para producción."], ["python manage.py sync_fuel_prices", "Consulta y guarda precios."], ["python manage.py seed_demo_data", "Carga datos de ejemplo; revisar antes de usar en producción."]], [3.1, 3.6])

    heading(doc, "12 Despliegue en Render")
    para(doc, "render.yaml define un servicio web Python y un cron semanal. El build instala requisitos y ejecuta collectstatic. El inicio usa gunicorn flamia.wsgi:application, escucha PORT, configura tres workers y timeout de 120 segundos.")
    bullets(doc, ["Crear y asociar PostgreSQL; copiar su URL a DATABASE_URL.", "Definir SECRET_KEY fuerte, DEBUG=False, ALLOWED_HOSTS y CSRF_TRUSTED_ORIGINS.", "Definir ORS_API_KEY y GOOGLE_MAPS_API_KEY si se utilizarán.", "Aplicar migraciones como paso de despliegue o release; render.yaml no las ejecuta actualmente.", "Crear el superusuario mediante una shell segura.", "Verificar /accounts/login/, archivos estáticos, mapa y una planificación de prueba."])
    para(doc, "El cron sync-fuel-prices se ejecuta los lunes a las 12:00 UTC, equivalente a las 06:00 en Guatemala, y comparte DATABASE_URL y SECRET_KEY con el servicio.")

    heading(doc, "13 Operación y mantenimiento")
    heading(doc, "Precios de combustible", 2)
    para(doc, "FuelPrice funciona como singleton con pk=1. El scraper intenta fuentes externas, valida rangos y conserva una fuente. Si todas fallan, el comando debe reportar el problema; el administrador puede introducir precios manualmente desde la UI.")
    heading(doc, "Base de datos", 2)
    bullets(doc, ["Respaldar PostgreSQL antes de migraciones y cambios masivos.", "Probar migraciones en una copia y revisar operaciones reversibles.", "Monitorear crecimiento de TripEvent, Payment y pedidos.", "No editar subtotales manualmente: OrderItem.save los recalcula.", "Tras modificar conexiones o coordenadas, invalidar la caché o esperar hasta 120 segundos."])
    heading(doc, "Observabilidad", 2)
    bullets(doc, ["Registrar errores del servidor con contexto y un identificador de solicitud.", "Vigilar respuestas 403, 400 y 500, latencia de APIs y fallos de proveedores externos.", "Alertar por fallo del cron de combustible y por falta de copias de seguridad.", "Revisar periódicamente cuentas inactivas, permisos staff y roles."])

    heading(doc, "14 Pruebas")
    para(doc, "La suite existente cubre el camino mínimo de Dijkstra, la creación de un viaje por API y la validación de capacidad. Ejecútela con python manage.py test.")
    heading(doc, "Cobertura recomendada", 2)
    bullets(doc, ["A estrella, rutas desconectadas, conexiones dirigidas y caché.", "Transiciones válidas e inválidas de viaje y parada.", "Permisos para cada rol en todos los POST.", "Idempotencia de cobros por UUID y sincronización offline.", "Cálculos de venta, pagos, saldo y filtros de fecha.", "Colisiones de códigos PED y VIA.", "Validación de coordenadas, importes negativos y cantidades cero.", "Pruebas de navegador para formularios, mapa y service worker.", "Despliegue de prueba con PostgreSQL, collectstatic y proxy HTTPS."])

    heading(doc, "15 Riesgos y deuda técnica")
    table(doc, ["Prioridad", "Hallazgo", "Acción sugerida"], [
        ["Alta", "Fallback de SECRET_KEY y DEBUG=True por defecto.", "Fallar al iniciar en producción si faltan variables seguras."],
        ["Alta", "Autorización inconsistente en endpoints de escritura.", "Definir matriz de permisos y aplicarla centralmente."],
        ["Alta", "render.yaml no aplica migraciones.", "Añadir proceso de release controlado."],
        ["Media", "Validación same_route anulada por True or.", "Decidir requisito y corregir o eliminar la condición."],
        ["Media", "Excepciones internas pueden exponerse al cliente.", "Usar logging y respuestas sanitizadas."],
        ["Media", "Planificación multi parada es voraz.", "Documentar alcance o incorporar optimización global."],
        ["Media", "Cobertura automatizada limitada.", "Priorizar permisos, finanzas y estados."],
        ["Baja", "Admin no registra Driver ni UserProfile.", "Añadirlos si el soporte debe gestionarlos desde /admin/."],
        ["Baja", "Comentarios y textos muestran inconsistencias de codificación en algunos archivos.", "Normalizar el repositorio a UTF-8 y añadir revisión."],
    ], [0.8, 3.0, 3.0])

    heading(doc, "16 Referencia rápida")
    table(doc, ["Elemento", "Valor"], [["Proyecto Django", "flamia"], ["Aplicación", "logistics"], ["Zona horaria", "America/Guatemala"], ["Idioma", "es-gt"], ["Login", "/accounts/login/"], ["Administración", "/admin/"], ["API", "/api/"], ["Base local", "db.sqlite3"], ["Base producción", "PostgreSQL mediante DATABASE_URL"], ["Caché de rutas", "120 segundos"], ["Service worker", "/service-worker.js"], ["Manifiesto", "/manifest.json"]], [2.2, 4.5])
    heading(doc, "Lista de entrega a producción", 2)
    bullets(doc, ["Pruebas aprobadas.", "Migraciones respaldadas y aplicadas.", "SECRET_KEY, hosts, CSRF y DEBUG revisados.", "Base de datos y copias de seguridad verificadas.", "Estáticos y login comprobados.", "Roles y cuentas revisados.", "Proveedores de mapas probados.", "Cron de combustible observado.", "Plan de reversión documentado."])
    doc.save(OUT / "Documentacion_Tecnica_Flamia.docx")


if __name__ == "__main__":
    make_diagrams()
    build_user_manual()
    build_technical()
    print(OUT / "Manual_de_Usuario_Flamia.docx")
    print(OUT / "Documentacion_Tecnica_Flamia.docx")
